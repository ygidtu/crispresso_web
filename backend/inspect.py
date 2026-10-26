"""Excel 预览与字段探测（不改动流水线逻辑，仅辅助前端预填配置）。"""
import openpyxl
import re

# yw4 解析器固定要求的列（config 步取 样品名称/B5/B7 引物；amplicons 步取 样品名称/ID/spacer/产物）。
REQUIRED_COLS = ["样品名称", "ID", "B5", "B7", "spacer", "产物（no adapter）"]
_DNA_RE = re.compile(r"^[ACGTN]+$", re.I)


def _iter_rows(ws, min_row, max_row):
    return list(ws.iter_rows(min_row=min_row, max_row=max_row, values_only=True))


def _has_primer(ws, col_idx, headers, sample_rows=20):
    """B5/B7 列本身或紧邻下一列，任一单元格含 ACGT 序列即视为有引物。"""
    for offset in (0, 1):
        idx = col_idx + offset
        if idx >= len(headers):
            continue
        for row in _iter_rows(ws, 2, 2 + sample_rows - 1):
            if idx < len(row) and row[idx] is not None:
                if _DNA_RE.match(str(row[idx]).strip().upper()):
                    return True
    return False


def check_sheet(xlsx_path, name):
    """校验单个 sheet 是否满足 yw4 parser 的列结构要求。"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[name]
        headers = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = ["" if c is None else str(c).strip() for c in row]
                continue
            break
        errors = []
        missing = [c for c in REQUIRED_COLS if c not in headers]
        if missing:
            errors.append("缺少必需列：" + "、".join(missing))
        for tag in ("B5", "B7"):
            if tag in headers:
                if not _has_primer(ws, headers.index(tag), headers):
                    errors.append(f"{tag} 未检测到引物序列（{tag} 列或紧邻下一列需为 ACGT 碱基序列）")
        if not headers:
            errors.append("表格为空，缺少表头行")
        ok = not errors
    finally:
        wb.close()
    return {"name": name, "columns": headers, "ok": ok, "errors": errors}


def check_sheets(xlsx_path):
    """校验所有 sheet，返回 {name: {ok, columns, errors}}。"""
    out = {}
    for name in list_sheets(xlsx_path):
        out[name] = check_sheet(xlsx_path, name)
    return out


def list_sheets(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def sheet_preview(xlsx_path, sheet, nrows=4):
    """返回表头与预览行（全部转字符串，None -> ""）。"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        headers, data = [], []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = ["" if c is None else str(c) for c in row]
                continue
            if i > nrows:
                break
            data.append(["" if c is None else str(c) for c in row])
    finally:
        wb.close()
    return {"columns": headers, "rows": data}


def suggest_sheet(xlsx_path, preferred=("样本信息", "4个样本", "12个样本")):
    """优先返回预设名 sheet，其次返回含 '样品名称' 列的 sheet。"""
    names = list_sheets(xlsx_path)
    for name in preferred:
        if name in names:
            return name
    for name in names:
        preview = sheet_preview(xlsx_path, name, nrows=1)
        if "样品名称" in preview["columns"]:
            return name
    return names[0] if names else None
